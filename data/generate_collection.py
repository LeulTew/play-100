"""Rebuild Play 100 from the unmodified, user-provided Excel workbook.

Usage:
    python generate_collection.py --source "path-to-original.xlsx"
    python generate_collection.py --source "path-to-original.xlsx" --inspect-only

Outputs are written beside this script unless --output is specified. No network
requests, source edits, macros, external workbook links, or review-score lookups
are performed.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import posixpath
import re
import sys
import textwrap
import unicodedata
import warnings
from collections import Counter
from fractions import Fraction
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import PIL
import xlsxwriter
from PIL import Image, ImageDraw, ImageFont, ImageOps


SOURCE_NAME = "AAA_games_u_have_to_play_list_top_100.xlsx"
PRIMARY_SHEET = "AAA Top 50"
REVIEWED_SOURCE_SHA256 = "301043fb513209b129586f617f1d205526c3c61785c48182760cf1aacaf29143"
CRITIC_KEYS = ("metacritic", "metacriticPc", "ign", "gamespot", "pcGamer")
CRITIC_MAX = dict(zip(CRITIC_KEYS, (100, 100, 10, 10, 100)))
WORKBOOK_NAME = "Play-100-Collection.xlsx"
TABLE_HEADER_ROW = 5
FIRST_DATA_ROW = TABLE_HEADER_ROW + 1
MAIN_HEADERS = (
    "Rank", "Game", "Tier", "Year", "Publisher / Studio", "Genre",
    "Metacritic /100", "Metacritic PC /100", "IGN /10", "GameSpot /10",
    "PC Gamer /100", "Critic average /100", "Leul's original rating /10", "Source note",
    "Why it ranks here", "Cover",
)
COLORS = {
    "background": "#F8F7F0", "panel": "#F3F4ED", "alternate": "#EAEDE3",
    "line": "#CBD3C3", "text": "#1D2921", "muted": "#536355",
    "secondary": "#3B4B3D", "ink": "#18211C", "onInk": "#F8F7F0",
    "onInkMuted": "#C8D2C4", "accent": "#D3EA82",
    "accentInk": "#354919", "accentPanel": "#EBF2D6",
}
NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
    "x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def write_json(path: Path, data: object) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def contrast_ratio(foreground: str, background: str) -> float:
    def luminance(color: str) -> float:
        channels = [int(color.lstrip("#")[index:index + 2], 16) / 255 for index in (0, 2, 4)]
        linear = [value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4 for value in channels]
        return sum(value * weight for value, weight in zip(linear, (0.2126, 0.7152, 0.0722)))
    first, second = sorted((luminance(foreground), luminance(background)))
    return (second + 0.05) / (first + 0.05)


def slugify(title: str) -> str:
    title = re.sub(r"['’]", "", title)
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_title.lower()).strip("-")


def critic_average(critics: dict) -> float | None:
    normalized = [
        value * (100 / CRITIC_MAX[key])
        for key, value in critics.items()
        if value is not None
    ]
    return math.fsum(normalized) / len(normalized) if normalized else None


def rank_index(rank: int) -> float:
    return float(Fraction(10) - Fraction((rank - 1) * 3, 99))


def extract_source_note(formula: object, cached: object) -> str | None:
    if isinstance(cached, str):
        suffix = re.sub(r"^\s*[+-]?\d+(?:\.\d+)?\s*", "", cached).strip()
        if suffix:
            return suffix
    if isinstance(formula, str) and not formula.startswith("="):
        suffix = re.sub(r"^\s*[+-]?\d+(?:\.\d+)?\s*", "", formula).strip()
        return suffix or None
    if isinstance(formula, str):
        match = re.search(r'&"([^"]+)"\s*$', formula)
        if match:
            return match.group(1).strip() or None
    return None


def original_rating(cached: object, raw: str, cell: str, number_format: str) -> dict:
    if is_number(cached):
        value = cached
        decimals = re.fullmatch(r"0\.(0+)", number_format)
        display = f"{value:.{len(decimals.group(1))}f}" if decimals else raw
        source_type = "number"
    elif isinstance(cached, str):
        match = re.match(r"^\s*([+-]?\d+(?:\.\d+)?)(?=\s|$)", cached)
        if not match:
            raise ValueError(f"Missing original author rating at {cell}")
        value = float(match.group(1))
        display = cached
        source_type = "text"
    else:
        raise ValueError(f"Missing original author rating cache at {cell}")
    if not 0 <= value <= 10:
        raise ValueError(f"Original author rating outside /10 scale at {cell}")
    return {
        "value": value, "rawValue": raw, "display": display,
        "sourceCell": cell, "numberFormat": number_format, "sourceType": source_type,
    }


def read_author_config(path: Path | None) -> dict:
    candidates = [path] if path is not None else [
        Path(__file__).resolve().parent.parent / "author.json",
        Path(__file__).resolve().with_name("author.json"),
    ]
    resolved = next((candidate for candidate in candidates if candidate.is_file()), None)
    if resolved is None:
        raise ValueError("Provide the shared author.json with --author-config.")
    author = json.loads(resolved.read_text(encoding="utf-8"))
    required = ("fullName", "shortName", "githubOwner", "repository", "githubUrl", "linkedinUrl", "telegramHandle", "telegramUrl", "siteUrl")
    if any(not isinstance(author.get(key), str) or not author[key].strip() for key in required):
        raise ValueError("Author configuration is incomplete.")
    if any(not author[key].startswith("https://") for key in ("githubUrl", "linkedinUrl", "telegramUrl", "siteUrl")):
        raise ValueError("Public author links must use HTTPS.")
    return author


def read_source(source: Path) -> tuple[dict, dict, list[dict]]:
    with ZipFile(source) as archive:
        core_properties = ET.fromstring(archive.read("docProps/core.xml"))
        created_node = core_properties.find("{http://purl.org/dc/terms/}created")
        source_document_created = created_node.text if created_node is not None else None
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = relationship_map(archive, "xl/workbook.xml")
        primary_xml = next(
            relations[sheet.attrib[f"{{{NS['r']}}}id"]]
            for sheet in workbook_xml.findall("m:sheets/m:sheet", NS)
            if sheet.attrib["name"] == PRIMARY_SHEET
        )
        source_cells = {
            cell.attrib["r"]: cell
            for cell in ET.fromstring(archive.read(primary_xml)).findall("m:sheetData/m:row/m:c", NS)
        }
        shared_strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_strings = [
                "".join(node.text or "" for node in item.findall(".//m:t", NS))
                for item in ET.fromstring(archive.read("xl/sharedStrings.xml")).findall("m:si", NS)
            ]
    with warnings.catch_warnings(record=True) as notices:
        warnings.simplefilter("always", UserWarning)
        formulas = openpyxl.load_workbook(source, data_only=False)
        values = openpyxl.load_workbook(source, data_only=True)
    primary = values[PRIMARY_SHEET]
    source_rows = {}
    records = []
    preserved_rows = []
    ignored_rows = []
    index_anomalies = []
    average_differences = []
    for row in primary.iter_rows(min_row=5):
        source_rank = row[0].value
        if not is_number(source_rank):
            nonempty = {cell.coordinate: cell.value for cell in row if cell.value is not None}
            if nonempty:
                ignored_rows.append(nonempty)
            continue
        if int(source_rank) != source_rank:
            raise ValueError(f"Non-integral rank at {row[0].coordinate}")
        rank = int(source_rank)
        if rank in source_rows:
            raise ValueError(f"Duplicate primary rank: {rank}")
        source_rows[rank] = row[0].row
        for column in (1, 3, 4, 12):
            if not isinstance(row[column].value, str) or not row[column].value.strip():
                raise ValueError(f"Missing required text at {row[column].coordinate}")
        if not is_number(row[2].value) or int(row[2].value) != row[2].value:
            raise ValueError(f"Invalid source year at {row[2].coordinate}")
        critics = {}
        for key, cell in zip(CRITIC_KEYS, row[5:10]):
            value = cell.value
            if value == "":
                value = None
            if value is not None and (not is_number(value) or not 0 <= value <= CRITIC_MAX[key]):
                raise ValueError(f"Invalid native-scale score at {cell.coordinate}: {value!r}")
            critics[key] = value
        original_formula = formulas[PRIMARY_SHEET].cell(row[0].row, 12).value
        note = extract_source_note(original_formula, row[11].value)
        rating_xml = source_cells[row[11].coordinate]
        raw_node = rating_xml.find("m:v", NS)
        if rating_xml.get("t") == "inlineStr":
            raw_rating = "".join(node.text or "" for node in rating_xml.findall(".//m:t", NS))
        elif raw_node is None or raw_node.text is None:
            raise ValueError(f"No original cached author rating at {row[11].coordinate}")
        elif rating_xml.get("t") == "s":
            raw_rating = shared_strings[int(raw_node.text)]
        else:
            raw_rating = raw_node.text
        author_rating = original_rating(row[11].value, raw_rating, row[11].coordinate, row[11].number_format)
        record = {
            "rank": rank,
            "slug": slugify(row[1].value),
            "title": row[1].value,
            "year": int(row[2].value),
            "studio": row[3].value,
            "genre": row[4].value,
            "genreTags": [part.strip() for part in row[4].value.split("/") if part.strip()],
            "tier": "core" if rank <= 50 else "essential",
            "critics": critics,
            "criticAverage": critic_average(critics),
            "rankIndex": rank_index(rank),
            "authorRating": author_rating,
            "rationale": row[12].value,
            "sourceNote": note,
            "artwork": None,
        }
        records.append(record)
        preserved_rows.append({
            "rank": rank,
            "sourceRow": row[0].row,
            "sourceCells": {
                cell.coordinate: {
                    "entered": formulas[PRIMARY_SHEET][cell.coordinate].value,
                    "cached": cell.value,
                    **({"cachedRaw": raw_rating, "numberFormat": cell.number_format} if cell.column == 12 else {}),
                }
                for cell in row[:13]
            },
        })
        if is_number(row[10].value):
            if not math.isclose(row[10].value, record["criticAverage"], abs_tol=1e-12):
                average_differences.append({
                    "rank": rank, "sourceCache": row[10].value,
                    "recomputed": record["criticAverage"],
                })
        if isinstance(original_formula, str) and original_formula.startswith("="):
            if f"A{row[0].row}" not in original_formula:
                index_anomalies.append({
                    "rank": rank, "title": record["title"], "cell": row[11].coordinate,
                    "sourceFormula": original_formula, "sourceCache": row[11].value,
                    "correctIndex": record["rankIndex"],
                    "handling": "The original cached author value is preserved. The separate legacy rankIndex field is not the author's recorded rating.",
                })
    records.sort(key=lambda item: item["rank"])
    if [item["rank"] for item in records] != list(range(1, 101)):
        raise ValueError("The primary sheet must contain exactly the unique ranks 1 through 100.")
    for key in ("title", "slug"):
        if len({item[key] for item in records}) != 100:
            raise ValueError(f"Primary {key} values are not unique.")
    methodology = values["Methodology"]
    collection = {
        "schemaVersion": 1,
        "collection": {
            "title": "Play 100",
            "sourceFile": SOURCE_NAME,
            "scope": methodology["B3"].value,
            "rankingBasis": methodology["B4"].value + " " + methodology["B5"].value,
            "criticScoresAreSnapshot": True,
        },
        "games": records,
    }
    by_title = {record["title"]: record for record in records}
    by_rank = {record["rank"]: record for record in records}
    timeline_entries = []
    for first in (1, 5):
        for row in range(4, values["Timeline"].max_row + 1):
            sheet = values["Timeline"]
            title = sheet.cell(row, first + 2).value
            if not title:
                continue
            game = by_title.get(title)
            timeline_entries.append({
                "cell": sheet.cell(row, first + 2).coordinate,
                "title": title,
                "sourceRank": sheet.cell(row, first + 1).value,
                "sourceYear": sheet.cell(row, first).value,
                "sourceTier": sheet.cell(row, first + 3).value,
                "canonicalRank": game["rank"] if game else None,
                "canonicalYear": game["year"] if game else None,
                "canonicalTier": game["tier"] if game else None,
            })
    timeline_titles = {item["title"] for item in timeline_entries}
    sortable_differences = []
    blank_to_zero = []
    for row in values["Sortable Ratings"].iter_rows(min_row=2):
        rank = row[0].value
        if rank not in source_rows:
            raise ValueError(f"Unexpected Sortable Ratings rank at {row[0].coordinate}")
        main_row = primary[source_rows[rank]]
        for cell, original_column in zip(row, (0, 1, 5, 6, 7, 8, 9, 10, 11)):
            original = main_row[original_column]
            if cell.value == original.value:
                continue
            entry = {
                "rank": rank, "title": by_rank[rank]["title"], "cell": cell.coordinate,
                "sourcePrimaryCell": original.coordinate,
                "derivativeCachedValue": cell.value, "authoritativeValue": original.value,
            }
            (blank_to_zero if original.value is None and cell.value == 0 else sortable_differences).append(entry)
    alternate_top_ten = []
    for row in range(4, 14):
        rank = methodology.cell(row, 4).value
        title = methodology.cell(row, 5).value
        alternate_top_ten.append({
            "rank": rank, "sourceTitle": title,
            "canonicalTitle": by_rank[rank]["title"],
            "matches": title == by_rank[rank]["title"],
        })
    audit = {
        "sourceFile": SOURCE_NAME,
        "sourceSha256": digest(source.read_bytes()),
        "sourceSizeBytes": source.stat().st_size,
        "sourceDocumentCreated": source_document_created,
        "primarySheet": PRIMARY_SHEET,
        "sourceScope": methodology["B3"].value,
        "sourceRanking": methodology["B4"].value,
        "sourceRankingSignals": methodology["B5"].value,
        "sourceTierLegend": primary["A2"].value,
        "libraryReadWarnings": sorted({str(notice.message) for notice in notices}),
        "libraryReadWarningHandling": (
            "Original is read only. Unsupported source formatting extensions are not "
            "round-tripped; the enhanced workbook is created independently."
        ),
        "sheets": [{
            "name": sheet.title,
            "dimensions": sheet.calculate_dimension(),
            "nonemptyCellCount": sum(cell.value is not None for row in sheet for cell in row),
            "formulaCount": sum(cell.data_type == "f" for row in sheet for cell in row),
            "embeddedImageCount": len(sheet._images),
            "mergedRanges": sorted(str(item) for item in sheet.merged_cells.ranges),
            "comments": {
                cell.coordinate: cell.comment.text
                for row in sheet for cell in row if cell.comment is not None
            },
            "hiddenRows": [index for index, item in sheet.row_dimensions.items() if item.hidden],
            "hiddenColumns": [index for index, item in sheet.column_dimensions.items() if item.hidden],
        } for sheet in formulas],
        "primaryRecords": preserved_rows,
        "ignoredNonRecordRows": ignored_rows,
        "extraColumnsNO": {
            cell.coordinate: cell.value
            for row in primary.iter_rows(min_col=14)
            for cell in row if cell.value is not None
        },
        "missingCriticCounts": {
            key: sum(game["critics"][key] is None for game in records) for key in CRITIC_KEYS
        },
        "sourceAverageCacheDifferences": average_differences,
        "rankIndexSourceAnomalies": index_anomalies,
        "sourceNotes": [{
            "rank": game["rank"], "title": game["title"], "sourceNote": game["sourceNote"],
        } for game in records if game["sourceNote"]],
        "methodologyAlternateTopTen": alternate_top_ten,
        "sortableRatings": {
            "recordCount": values["Sortable Ratings"].max_row - 1,
            "blankPrimaryScoresDisplayedAsZero": blank_to_zero,
            "otherDifferences": sortable_differences,
        },
        "timeline": {
            "recordCount": len(timeline_entries),
            "uniqueTitleCount": len(timeline_titles),
            "missingYearCount": sum(item["sourceYear"] is None for item in timeline_entries),
            "missingTierCount": sum(item["sourceTier"] is None for item in timeline_entries),
            "entries": timeline_entries,
            "titlesNotInPrimary": [item["title"] for item in timeline_entries if item["title"] not in by_title],
            "primaryTitlesOmitted": [game["title"] for game in records if game["title"] not in timeline_titles],
            "rankDisagreementCount": sum(
                item["sourceRank"] != item["canonicalRank"]
                for item in timeline_entries if item["canonicalRank"] is not None
            ),
        },
    }
    candidates = read_embedded_images(source, values, by_title, by_rank, source_rows)
    audit["embeddedImages"] = [public_image(candidate) for candidate in candidates]
    formulas.close()
    values.close()
    return collection, audit, candidates


def relationship_map(archive: ZipFile, owner: str) -> dict[str, str]:
    directory, filename = posixpath.split(owner)
    rel_path = posixpath.join(directory, "_rels", filename + ".rels")
    if rel_path not in archive.namelist():
        return {}
    relationships = ET.fromstring(archive.read(rel_path))
    result = {}
    for relationship in relationships:
        if relationship.get("TargetMode") == "External":
            continue
        target = relationship.attrib["Target"]
        result[relationship.attrib["Id"]] = (
            target.lstrip("/") if target.startswith("/")
            else posixpath.normpath(posixpath.join(directory, target))
        )
    return result


def read_embedded_images(
    source: Path, values, by_title: dict, by_rank: dict, source_rows: dict,
) -> list[dict]:
    candidates = []
    row_ranks = {row: rank for rank, row in source_rows.items()}
    with ZipFile(source) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook.find("m:sheets", NS)
        workbook_links = relationship_map(archive, "xl/workbook.xml")
        for sheet in sheets:
            sheet_name = sheet.attrib["name"]
            sheet_path = workbook_links[sheet.attrib[f"{{{NS['r']}}}id"]]
            sheet_xml = ET.fromstring(archive.read(sheet_path))
            sheet_links = relationship_map(archive, sheet_path)
            for drawing in sheet_xml.findall("m:drawing", NS):
                drawing_path = sheet_links[drawing.attrib[f"{{{NS['r']}}}id"]]
                image_links = relationship_map(archive, drawing_path)
                for anchor in ET.fromstring(archive.read(drawing_path)):
                    start = anchor.find("x:from", NS)
                    blip = anchor.find(".//a:blip", NS)
                    if start is None or blip is None:
                        raise ValueError(f"Unsupported image anchor in {drawing_path}")
                    row = int(start.find("x:row", NS).text) + 1
                    column = int(start.find("x:col", NS).text) + 1
                    if sheet_name == PRIMARY_SHEET and row in row_ranks and column == 2:
                        game = by_rank[row_ranks[row]]
                        label_cell = f"B{row}"
                    elif sheet_name == "Posters":
                        label = values[sheet_name].cell(row - 1, column)
                        game = by_title.get(label.value)
                        label_cell = label.coordinate
                        if game is None:
                            raise ValueError(f"Poster label has no exact title match: {label.value!r}")
                    else:
                        raise ValueError(f"Unmapped image at {sheet_name}!{row},{column}")
                    media_path = image_links[blip.attrib[f"{{{NS['r']}}}embed"]]
                    payload = archive.read(media_path)
                    with Image.open(io.BytesIO(payload)) as image:
                        image.verify()
                    with Image.open(io.BytesIO(payload)) as image:
                        size, image_format = image.size, image.format
                    candidates.append({
                        "rank": game["rank"], "title": game["title"], "slug": game["slug"],
                        "sourceSheet": sheet_name, "sourceLabelCell": label_cell,
                        "sourceAnchor": {"row": row, "column": column},
                        "sourceMediaPath": media_path, "sha256": digest(payload),
                        "width": size[0], "height": size[1], "format": image_format,
                        "mappingBasis": "Exact title at the source row or immediately above the poster anchor.",
                        "_payload": payload,
                    })
    return candidates


def public_image(candidate: dict) -> dict:
    return {key: value for key, value in candidate.items() if not key.startswith("_")}


def make_contact_sheets(candidates: list[dict], output: Path) -> list[str]:
    preview = output / "previews"
    preview.mkdir(exist_ok=True)
    paths = []
    groups = [
        ("source-thumbnails", [item for item in candidates if item["sourceSheet"] == PRIMARY_SHEET], 25),
        ("source-posters", [item for item in candidates if item["sourceSheet"] == "Posters"], 9),
    ]
    font = ImageFont.load_default(size=15)
    for prefix, group, per_page in groups:
        for offset in range(0, len(group), per_page):
            items = group[offset:offset + per_page]
            columns = 5 if per_page == 25 else 3
            rows = math.ceil(len(items) / columns)
            cell_width, cell_height = 250, 286
            canvas = Image.new("RGB", (columns * cell_width, rows * cell_height + 45), COLORS["ink"])
            draw = ImageDraw.Draw(canvas)
            draw.text((16, 12), f"SOURCE IMAGE MAPPING / {prefix} / {offset + 1}-{offset + len(items)}", font=font, fill=COLORS["accent"])
            for index, candidate in enumerate(items):
                left = (index % columns) * cell_width
                top = (index // columns) * cell_height + 45
                with Image.open(io.BytesIO(candidate["_payload"])) as image:
                    cover = ImageOps.contain(image.convert("RGB"), (172, 204), Image.Resampling.LANCZOS)
                canvas.paste(cover, (left + (cell_width - cover.width) // 2, top + 3))
                label = f'{candidate["rank"]:03d}  {candidate["title"]}'
                for line_index, line in enumerate(textwrap.wrap(label, width=30)):
                    draw.text((left + 12, top + 215 + 18 * line_index), line, font=font, fill=COLORS["onInk"])
            filename = f"{prefix}-{offset // per_page + 1:02d}.jpg"
            canvas.save(preview / filename, quality=93, optimize=True)
            paths.append("previews/" + filename)
    return paths


def select_artwork(collection: dict, candidates: list[dict], output: Path) -> list[dict]:
    assets = output / "assets"
    assets.mkdir(exist_ok=True)
    selected = []
    for game in collection["games"]:
        options = [candidate for candidate in candidates if candidate["rank"] == game["rank"]]
        if not options:
            continue
        choice = max(
            options,
            key=lambda item: (item["width"] * item["height"], item["sourceSheet"] == "Posters"),
        )
        if choice["title"] != game["title"]:
            raise ValueError(f"Artwork title mismatch for rank {game['rank']}")
        extension = {"JPEG": ".jpg", "PNG": ".png"}.get(choice["format"])
        if extension is None:
            raise ValueError(f"Unsupported reusable image format: {choice['format']}")
        relative = "assets/" + game["slug"] + extension
        (assets / (game["slug"] + extension)).write_bytes(choice["_payload"])
        game["artwork"] = {"file": relative, "source": "User-provided workbook"}
        selected.append({
            **public_image(choice), "file": relative,
            "selection": "Largest supplied image for this exact anchored title; original bytes preserved.",
            "review": "Source anchors, labels, and visible cover branding reviewed together.",
        })
    return selected


def average_formula(excel_row: int) -> str:
    return (
        f'=IF(COUNT(G{excel_row}:K{excel_row})=0,"",'
        f'(SUM(G{excel_row}:H{excel_row},K{excel_row})+10*SUM(I{excel_row}:J{excel_row}))'
        f'/COUNT(G{excel_row}:K{excel_row}))'
    )


def index_formula(excel_row: int) -> str:
    return f"=10-((A{excel_row}-1)*(3/99))"


def main_lookup_formula(column: str, rank_reference: str) -> str:
    return (
        f'=IFERROR(INDEX(\'The 100\'!${column}$7:${column}$106,'
        f'MATCH({rank_reference},\'The 100\'!$A$7:$A$106,0)),"")'
    )


def main_link_formula(rank_reference: str, label: str = "Open entry") -> str:
    label = label.replace('"', '""')
    return (
        '=HYPERLINK("#\'The 100\'!B"&'
        f'(MATCH({rank_reference},\'The 100\'!$A$7:$A$106,0)+6),"{label}")'
    )


def omit_generated_document_timestamps(path: Path) -> None:
    # Spreadsheet libraries invent creation dates when the source has none.
    # Omit these optional package fields rather than treating them as provenance.
    rebuilt = io.BytesIO()
    with ZipFile(io.BytesIO(path.read_bytes())) as original, ZipFile(rebuilt, "w") as target:
        for entry in original.infolist():
            payload = original.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                root = ET.fromstring(payload)
                for local_name in ("created", "modified"):
                    node = root.find("{http://purl.org/dc/terms/}" + local_name)
                    if node is not None:
                        root.remove(node)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, payload)
    path.write_bytes(rebuilt.getvalue())


def build_workbook(collection: dict, audit: dict, selected: list[dict], output: Path) -> None:
    games = collection["games"]
    author = collection["collection"]["author"]
    by_rank = {game["rank"]: game for game in games}
    images = {item["rank"]: item for item in selected}
    path = output / WORKBOOK_NAME
    workbook = xlsxwriter.Workbook(path, {
        "strings_to_formulas": False, "strings_to_urls": False,
        "use_zip64": False, "in_memory": True,
    })
    properties = {
        "title": "Play 100 | The gaming collection",
        "subject": "User-curated ranking and source critic-score snapshot",
        "author": author["fullName"],
        "comments": "Scores, years and ranking are source snapshots, not independently reverified.",
    }
    workbook.set_properties(properties)
    workbook.set_calc_mode("auto")

    def fmt(**overrides):
        return workbook.add_format({
            "font_name": "Aptos", "font_size": 11, "font_color": COLORS["text"],
            "bg_color": COLORS["background"], "valign": "vcenter",
            **overrides,
        })

    base = fmt()
    title_format = fmt(font_size=32, bold=True, bg_color=COLORS["ink"], font_color=COLORS["onInk"])
    subtitle = fmt(font_size=11, font_color=COLORS["muted"], text_wrap=True)
    section = fmt(font_size=12, bold=True, font_color=COLORS["accentInk"])
    nav = fmt(font_size=11, bold=True, font_color=COLORS["accentInk"], underline=1)
    small = fmt(font_size=10, font_color=COLORS["muted"], text_wrap=True)
    header = fmt(
        font_size=10, bold=True, bg_color=COLORS["ink"], font_color=COLORS["onInk"], text_wrap=True,
        bottom=2, bottom_color=COLORS["accent"],
    )
    standard_link = fmt(
        font_size=10, font_color=COLORS["accentInk"], underline=1,
        bg_color=COLORS["panel"], text_wrap=True,
    )
    body_formats = {}
    author_formats = {}
    for stripe in (0, 1):
        background = COLORS["panel"] if stripe == 0 else COLORS["alternate"]
        common = {
            "bg_color": background, "text_wrap": True,
            "bottom": 1, "bottom_color": COLORS["line"],
        }
        body_formats[stripe] = {
            "text": fmt(**common),
            "game": fmt(**common, bold=True),
            "rank": fmt(**common, font_size=14, bold=True, align="center", font_color=COLORS["accentInk"], num_format="0"),
            "year": fmt(**common, align="center", num_format="0"),
            "score": fmt(**common, align="center", num_format="0.0"),
            "core": fmt(**common, align="center", font_size=10, bold=True, font_color=COLORS["accentInk"]),
            "essential": fmt(**common, align="center", font_size=10, font_color=COLORS["secondary"]),
            "note": fmt(**common, font_size=10, font_color=COLORS["accentInk"]),
            "link": fmt(**common, font_size=10, font_color=COLORS["accentInk"], underline=1),
        }
        for number_format in sorted({game["authorRating"]["numberFormat"] for game in games}):
            author_formats[(stripe, number_format)] = fmt(**common, align="center", num_format=number_format)
    start = workbook.add_worksheet("Start Here")
    main = workbook.add_worksheet("The 100")
    chronology = workbook.add_worksheet("Chronology")
    gallery = workbook.add_worksheet("Cover Gallery")
    methodology = workbook.add_worksheet("Methodology")
    for sheet in (start, main, chronology, gallery, methodology):
        sheet.hide_gridlines(2)
        sheet.set_default_row(21)
        sheet.set_tab_color(COLORS["accent"] if sheet is main else COLORS["ink"])
        sheet.set_margins(0.3, 0.3, 0.4, 0.4)
        sheet.set_header("&LPlay 100&RUser-curated collection", {"margin": 0.15})
        sheet.set_footer(
            f'&L&8Curated by {author["fullName"]}\n{author["githubUrl"]}'
            f'&C&8{author["linkedinUrl"]}\n{author["telegramUrl"]}'
            '&R&8&P / &N', {"margin": 0.15},
        )

    def attribution(sheet, row: int, last_column: int):
        sheet.set_row(row, 25)
        sheet.merge_range(row, 0, row, last_column, f'Curated by {author["fullName"]}', section)
        boundaries = (0, (last_column + 1) // 3, 2 * (last_column + 1) // 3, last_column + 1)
        links = (
            (author["githubUrl"], f'GitHub · {author["githubOwner"]}/{author["repository"]}'),
            (author["linkedinUrl"], f'LinkedIn · {author["fullName"]}'),
            (author["telegramUrl"], f'Telegram · {author["telegramHandle"]}'),
        )
        sheet.set_row(row + 1, 32)
        for index, (url, label) in enumerate(links):
            first, last = boundaries[index], boundaries[index + 1] - 1
            sheet.merge_range(row + 1, first, row + 1, last, "", standard_link)
            sheet.write_url(row + 1, first, url, standard_link, label)
        sheet.print_area(0, 0, row + 2, last_column)

    def fill_canvas(sheet, last_row: int, last_column: int):
        for row in range(last_row + 1):
            for column in range(last_column + 1):
                sheet.write_blank(row, column, None, base)

    def banner(sheet, heading: str, description: str, last_column: int):
        sheet.set_row(0, 44)
        sheet.set_row(1, 31)
        sheet.merge_range(0, 0, 0, last_column, heading, title_format)
        sheet.merge_range(1, 0, 1, last_column, description, subtitle)
        sheet.write_url(3, 0, "internal:'Start Here'!A1", nav, "← Start Here")

    def table(sheet, name: str, headers: tuple, count: int, first_row: int = TABLE_HEADER_ROW):
        sheet.add_table(first_row, 0, first_row + count, len(headers) - 1, {
            "name": name, "style": None, "banded_rows": False,
            "columns": [{"header": name, "header_format": header} for name in headers],
        })
        sheet.set_row(first_row, 40)
        sheet.repeat_rows(first_row, first_row)
        sheet.print_area(0, 0, first_row + count, len(headers) - 1)
        sheet.set_landscape()
        sheet.set_paper(8)
        sheet.fit_to_pages(1, 0)

    main.set_zoom(85)
    widths = (7, 40, 14, 9, 33, 30, 14, 16, 11, 13, 14, 17, 15, 26, 72, 12)
    for column, width in enumerate(widths):
        main.set_column(column, column, width, base)
    fill_canvas(main, 106, 15)
    banner(
        main, "THE 100",
        "Curated order, not a critic leaderboard. Core = ranks 1–50 · Essential = ranks 51–100 · Nintendo excluded.",
        15,
    )
    main.write_url(3, 2, "internal:'Chronology'!A1", nav, "Chronology →")
    main.write_url(3, 5, "internal:'Cover Gallery'!A1", nav, "Covers →")
    main.write_url(3, 8, "internal:'Methodology'!A1", nav, "Methodology →")
    main.merge_range(
        4, 0, 4, 15,
        "Leul's original rating is copied from his original rank-based column, including rounded/text caches. Critic snapshots stay separate. Rank + Game stay frozen.",
        small,
    )
    main.freeze_panes(FIRST_DATA_ROW, 2)
    main.set_selection(FIRST_DATA_ROW, 1, FIRST_DATA_ROW, 1)
    for game in games:
        row = TABLE_HEADER_ROW + game["rank"]
        excel_row = row + 1
        formats = body_formats[(game["rank"] - 1) % 2]
        main.set_row(row, 64)
        main.write_number(row, 0, game["rank"], formats["rank"])
        main.write_string(row, 1, game["title"], formats["game"])
        main.write_string(row, 2, game["tier"].title(), formats[game["tier"]])
        main.write_number(row, 3, game["year"], formats["year"])
        main.write_string(row, 4, game["studio"], formats["text"])
        main.write_string(row, 5, game["genre"], formats["text"])
        for column, key in enumerate(CRITIC_KEYS, start=6):
            value = game["critics"][key]
            if value is None:
                main.write_blank(row, column, None, formats["score"])
            else:
                main.write_number(row, column, value, formats["score"])
        main.write_formula(row, 11, average_formula(excel_row), formats["score"], game["criticAverage"] if game["criticAverage"] is not None else "")
        original = game["authorRating"]
        main.write_number(row, 12, original["value"], author_formats[((game["rank"] - 1) % 2, original["numberFormat"])])
        main.write_comment(row, 12, f'{author["fullName"]} — original rating\nSource: {PRIMARY_SHEET}!{original["sourceCell"]}\nExact cached text: {original["rawValue"]}\nOriginal header: my rating(based on rank)\nPreserved, not recomputed. Source annotations are also shown in Source note.')
        if game["sourceNote"]:
            main.write_string(row, 13, game["sourceNote"], formats["note"])
        else:
            main.write_blank(row, 13, None, formats["note"])
        main.write_string(row, 14, game["rationale"], formats["text"])
        card_row = 6 + ((game["rank"] - 1) // 4) * 12
        card_column = ((game["rank"] - 1) % 4) * 5
        location = openpyxl.utils.get_column_letter(card_column + 1) + str(card_row + 1)
        if game["artwork"]:
            main.write_url(row, 15, f"internal:'Cover Gallery'!{location}", formats["link"], "View cover")
        else:
            main.write_blank(row, 15, None, formats["link"])
    table(main, "Play100", MAIN_HEADERS, 100)
    for first, last, maximum in ((6, 7, 100), (8, 9, 10), (10, 10, 100)):
        main.data_validation(FIRST_DATA_ROW, first, 105, last, {
            "validate": "decimal", "criteria": "between", "minimum": 0, "maximum": maximum,
            "ignore_blank": True, "error_type": "stop", "show_error": True,
            "error_title": "Keep the native scale",
            "error_message": f"Enter a number from 0 to {maximum}, or leave blank if unavailable.",
            "input_title": f"Source critic score /{maximum}",
            "input_message": "Blank means missing. Do not substitute zero for missing scores.",
        })
    main.conditional_format("L7:L106", {
        "type": "data_bar", "bar_color": COLORS["accent"],
        "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 100,
    })
    main.conditional_format("N7:N106", {
        "type": "no_blanks",
        "format": fmt(bg_color=COLORS["accentPanel"], font_color=COLORS["accentInk"], text_wrap=True, font_size=10),
    })
    main.conditional_format("A7:P106", {
        "type": "formula", "criteria": "=$A7=50",
        "format": workbook.add_format({"bottom": 2, "bottom_color": COLORS["accent"]}),
    })
    main.write_comment("L6", "Average of available numeric entries. IGN and GameSpot are multiplied by 10. Both entered Metacritic columns count, so this is not an average of independent publications.")
    main.write_comment("M6", f'{author["fullName"]}\'s original ratings from column L, headed "my rating(based on rank)". Original cached values are retained, including rounded text values; they are not replaced by a new formula. Visitor/device-only ratings are a separate website feature.')
    main.write_comment("N6", "Only explicit source annotations are preserved. Blank does not mean played.")
    main.set_h_pagebreaks([56])

    chronology.set_zoom(90)
    chronology_headers = (
        "Year", "Rank", "Game", "Tier", "Publisher / Studio", "Genre",
        "Critic average /100", "Leul's original rating /10", "Source note", "Entry",
    )
    for column, width in enumerate((9, 8, 43, 14, 34, 32, 19, 17, 26, 14)):
        chronology.set_column(column, column, width, base)
    fill_canvas(chronology, 106, 9)
    banner(
        chronology, "A CHRONOLOGY OF THE COLLECTION",
        "All 100 entries, ordered by the source year, then curated rank. These are source years—not independently verified release dates.",
        9,
    )
    chronology.write_url(3, 2, "internal:'The 100'!A1", nav, "The 100 →")
    chronology.write_url(3, 5, "internal:'Methodology'!A1", nav, "How to read →")
    chronology.merge_range(4, 0, 4, 9, "Ranks stay curated. Critics and Leul's preserved original ratings look up The 100 by rank; links remain correct after sorting.", small)
    chronology.freeze_panes(FIRST_DATA_ROW, 3)
    chronological_games = sorted(games, key=lambda game: (game["year"], game["rank"]))
    for position, game in enumerate(chronological_games):
        row = FIRST_DATA_ROW + position
        excel_row = row + 1
        formats = body_formats[position % 2]
        chronology.set_row(row, 45)
        chronology.write_number(row, 0, game["year"], formats["year"])
        chronology.write_number(row, 1, game["rank"], formats["rank"])
        chronology.write_string(row, 2, game["title"], formats["game"])
        chronology.write_string(row, 3, game["tier"].title(), formats[game["tier"]])
        chronology.write_string(row, 4, game["studio"], formats["text"])
        chronology.write_string(row, 5, game["genre"], formats["text"])
        chronology.write_formula(row, 6, main_lookup_formula("L", f"B{excel_row}"), formats["score"], game["criticAverage"] if game["criticAverage"] is not None else "")
        chronology.write_formula(row, 7, main_lookup_formula("M", f"B{excel_row}"), author_formats[(position % 2, game["authorRating"]["numberFormat"])], game["authorRating"]["value"])
        if game["sourceNote"]:
            chronology.write_string(row, 8, game["sourceNote"], formats["note"])
        else:
            chronology.write_blank(row, 8, None, formats["note"])
        chronology.write_formula(row, 9, main_link_formula(f"B{excel_row}"), formats["link"], "Open entry")
    table(chronology, "Play100Chronology", chronology_headers, 100)
    chronology.conditional_format("A7:J106", {
        "type": "formula", "criteria": "=AND(ISNUMBER($A6),$A7<>$A6)",
        "format": workbook.add_format({"top": 2, "top_color": "#89988A"}),
    })

    gallery.set_zoom(90)
    gallery.set_column(0, 18, 8, base)
    fill_canvas(gallery, 306, 18)
    banner(
        gallery, "THE COVER GALLERY",
        "100 supplied covers · curated order · nine larger Posters images preferred over duplicate thumbnails. Original image bytes are retained.",
        18,
    )
    gallery.write_url(3, 5, "internal:'The 100'!A1", nav, "The 100 →")
    gallery.write_url(3, 10, "internal:'Methodology'!A1", nav, "Artwork notes →")
    gallery.merge_range(4, 0, 4, 18, "Representative source artwork, not proof of platform availability or exact edition. Small originals may look soft. Rank 73 uses source art branded HITMAN III.", small)
    gallery.freeze_panes(6, 0)
    gallery_title = fmt(bg_color=COLORS["ink"], font_color=COLORS["onInk"], font_size=12, bold=True, text_wrap=True, valign="top")
    gallery_label = fmt(bg_color=COLORS["panel"], font_size=10, font_color=COLORS["muted"], text_wrap=True)
    card_fill = fmt(bg_color=COLORS["panel"])
    for game in games:
        row = 6 + ((game["rank"] - 1) // 4) * 12
        column = ((game["rank"] - 1) % 4) * 5
        gallery.merge_range(row, column, row + 2, column + 3, f'{game["rank"]:03d}  {game["title"]}', gallery_title)
        gallery.merge_range(row + 3, column, row + 8, column + 3, "", card_fill)
        gallery.merge_range(row + 9, column, row + 9, column + 3, f'{game["year"]}  /  {game["tier"].title()}', gallery_label)
        gallery.merge_range(row + 10, column, row + 10, column + 3, "", standard_link)
        gallery.write_formula(row + 10, column, main_link_formula(str(game["rank"]), "View this entry →"), standard_link, "View this entry →")
        if game["artwork"]:
            image = images[game["rank"]]
            scale = min(155 / image["width"], 159 / image["height"])
            rendered_width = image["width"] * scale
            x_offset = max(3, int((244 - rendered_width) / 2))
            gallery.insert_image(row + 3, column, str(output / Path(game["artwork"]["file"])), {
                "x_scale": scale, "y_scale": scale, "x_offset": x_offset, "y_offset": 2,
                "object_position": 1,
                "description": f'{game["title"]} — user-provided workbook cover',
            })
    gallery.set_landscape()
    gallery.set_paper(9)
    gallery.fit_to_pages(1, 0)
    gallery.repeat_rows(0, 4)
    gallery.set_h_pagebreaks([6 + page * 24 for page in range(1, 13)])
    gallery.print_area(0, 0, 305, 18)

    start.set_zoom(95)
    start.set_column(0, 15, 8.5, base)
    fill_canvas(start, 56, 15)
    hero_background = fmt(bg_color=COLORS["ink"], font_color=COLORS["onInk"])
    for row in range(10):
        for column in range(16):
            start.write_blank(row, column, None, hero_background)
    start.set_row(0, 41)
    start.set_row(1, 32)
    start.merge_range("A1:J2", "PLAY 100", fmt(font_size=44, bold=True, bg_color=COLORS["ink"], font_color=COLORS["onInk"]))
    start.merge_range("A3:J4", "A collection worth making time for.", fmt(font_size=19, bg_color=COLORS["ink"], font_color=COLORS["onInk"], text_wrap=True))
    start.merge_range("A6:J8", "100 essential gaming experiences, in your curated order.\nA core 50. A next-essential 50. One collection.", fmt(font_size=13, text_wrap=True, bg_color=COLORS["ink"], font_color=COLORS["onInkMuted"]))
    start.merge_range("K1:P9", "", hero_background)
    if 1 in images:
        image = images[1]
        scale = min(190 / image["width"], 225 / image["height"])
        start.insert_image("L1", str(output / Path(by_rank[1]["artwork"]["file"])), {
            "x_scale": scale, "y_scale": scale, "x_offset": 17, "y_offset": 12,
            "description": "Rank 1 — Red Dead Redemption 2, user-provided cover",
        })
    start.merge_range("K10:P10", "01 / RED DEAD REDEMPTION 2", fmt(font_size=10, bold=True, bg_color=COLORS["ink"], font_color=COLORS["accent"]))
    stats = [
        (0, "100", "CURATED ENTRIES"),
        (4, "50 + 50", "CORE + ESSENTIAL"),
        (8, f'{min(game["year"] for game in games)}–{max(game["year"] for game in games)}', "SOURCE YEARS"),
        (12, str(len(selected)), "SUPPLIED COVERS"),
    ]
    for column, value, label in stats:
        start.merge_range(11, column, 12, column + 3, value, fmt(font_size=25, bold=True, bg_color=COLORS["panel"], align="center"))
        start.merge_range(13, column, 13, column + 3, label, fmt(font_size=9, bold=True, bg_color=COLORS["panel"], font_color=COLORS["muted"], align="center"))
    start.merge_range("A16:P16", "EXPLORE THE COLLECTION", section)
    navigation = [
        (17, 0, 7, "The 100", "01  THE 100 →", "Filter by tier, year, studio, genre or score. Rank and Game stay frozen."),
        (17, 8, 15, "Chronology", "02  CHRONOLOGY →", "The full collection ordered by source year, with the curated rank intact."),
        (22, 0, 7, "Cover Gallery", "03  COVER GALLERY →", "Every supplied cover, with a link back to its entry. No downloaded replacements."),
        (22, 8, 15, "Methodology", "04  METHODOLOGY →", "What each number means, what remains unknown, and how the source was reconciled."),
    ]
    for row, column, end_column, target, label, description in navigation:
        start.merge_range(row, column, row + 1, end_column, "", fmt(bg_color=COLORS["ink"]))
        start.write_url(row, column, f"internal:'{target}'!A1", fmt(bg_color=COLORS["ink"], bold=True, font_size=13, font_color=COLORS["accent"], underline=1), label)
        start.merge_range(row + 2, column, row + 3, end_column, description, fmt(bg_color=COLORS["panel"], font_color=COLORS["muted"], text_wrap=True))
    start.merge_range("A29:P29", "READ THE NUMBERS CORRECTLY", section)
    legend = [
        ("A31:D34", "CURATED RANK", "The order is the source collection's must-play ranking—not a critic leaderboard."),
        ("E31:H34", "CRITIC SNAPSHOT", "Scores retain their native /100 or /10 scales. No current verification or score date is claimed."),
        ("I31:L34", "AVAILABLE, NOT ZERO", "Blank critic cells are unknown. Averages skip blanks and normalize /10 entries to /100."),
        ("M31:P34", "LEUL'S ORIGINAL RATING", "Leul's rating /10 comes from his original rank-based column. Exact cached values and source notes are preserved."),
    ]
    for cell_range, label, body in legend:
        start.merge_range(cell_range, label + "\n\n" + body, fmt(font_size=10, text_wrap=True, valign="top", font_color=COLORS["secondary"]))
    start.merge_range("A37:P39", "Both Metacritic columns count when entered. The critic average describes available score columns, not independent publications. It supports browsing; it never determines the order.", fmt(font_size=11, bg_color=COLORS["panel"], font_color=COLORS["muted"], text_wrap=True))
    start.merge_range("A42:P44", "Twelve original ratings explicitly carry Leul's source annotation “(AI – not played)”. The numeric rating and the exact note are both preserved. These historical author notes never set a website visitor's played status.", fmt(font_size=11, text_wrap=True, font_color=COLORS["muted"]))
    start.merge_range("A47:P49", collection["collection"]["scope"], fmt(font_size=10, text_wrap=True, font_color=COLORS["muted"]))
    start.merge_range("A52:P54", "SOURCE  /  " + SOURCE_NAME + "\nThe primary list is authoritative. Stale side tables were reconciled; native score entries were not rewritten. See Methodology and artifact-manifest.json for the complete provenance.", fmt(font_size=9, text_wrap=True, font_color=COLORS["muted"]))
    start.print_area("A1:P54")
    start.set_paper(9)
    start.fit_to_pages(1, 1)
    start.activate()
    start.set_first_sheet()

    methodology.set_zoom(95)
    methodology.set_column(0, 2, 11, base)
    methodology.set_column(3, 14, 8.5, base)
    notes = [
        ("01 / SCOPE", collection["collection"]["scope"]),
        ("02 / RANKING", audit["sourceRanking"] + " Signals in the source: " + audit["sourceRankingSignals"]),
        ("03 / TIERS", "Core = ranks 1–50. Essential = ranks 51–100. The division is editorial, not a score cutoff. No source title was added, removed, promoted or demoted."),
        ("04 / NATIVE SCALES", "Metacritic, Metacritic PC and PC Gamer use /100. IGN and GameSpot use /10. A blank means no numeric value was entered in the primary snapshot; it never means zero."),
        ("05 / CRITIC AVERAGE", "Normalize IGN and GameSpot by multiplying by 10; then average only numeric, available columns. Both Metacritic columns count separately when entered. This is not an average of independent publications and is not the ranking algorithm."),
        ("06 / LEUL'S ORIGINAL RATING", "The original column L is headed “my rating(based on rank)”. It contains Leul's own ratings, based on his curated ranking. The original cached values are copied exactly; rounded TEXT results remain rounded (for example The Witcher 3 = 9.9 and Grand Theft Auto IV = 9.8). No formula repairs or reconstructed curve replace those values. This is separate from critic averages and any visitor's private website rating."),
        ("07 / SOURCE NOTES", "Twelve author-rating cells contain “(AI – not played)”. Their numeric author rating and original note are both preserved. An absent source note does not establish play history. Historical author notes never set a visitor's Played/Completed state."),
        ("08 / WHAT IS VERIFIED", "The extraction, arithmetic, record count, source preservation, formula caches and image mapping were checked. Critic scores, game years, platform editions and release metadata were not independently reverified. No critic snapshot date is known."),
        ("09 / DATA PRESERVATION", "Game titles, studios, source years, genres, native critic entries and ranking rationales come from the primary sheet unchanged. Non-AAA premium exceptions already present in the source remain included. JSON genreTags only split the existing genre text at slash separators."),
        ("10 / FORMULAS", "The 100 recalculates critic averages from entered columns. Leul's original ratings are fixed copied numbers so recalculation cannot silently rewrite them. Chronology looks up both metrics by rank with matching initial caches. Sorting either table preserves associations. Formula mode is automatic; there are no macros or external workbook links."),
        ("11 / COVER PROVENANCE", "109 original embedded images were found: 100 row thumbnails and nine larger gallery posters. For each game, the largest correctly anchored supplied image is reused without altering its bytes. All nine poster labels match exact primary titles. No artwork was fetched, invented or replaced."),
        ("12 / COVER LIMITATIONS", "The source is low resolution: 91 selected covers are 120 pixels high; nine are larger. Covers are representative editions and cannot establish platform availability. Rank 73 is titled Hitman: World of Assassination with source year 2016, while its supplied cover is branded HITMAN III. That source naming/edition caveat is retained, not silently corrected."),
        ("13 / PRIMARY PRESENTATION", "A malformed average header was replaced with Critic average /100. Leul's original rating ownership is explicit. Rank 7's original cached 9.8 is retained despite its historical hard-coded formula; no author value is repaired or recalculated. The tier divider is not a game. Columns N and O contained no additional source data."),
        ("14 / DERIVATIVE REPAIRS", "Sortable Ratings had 100 entries but displayed all 56 missing primary critic values as zeros through direct-cell-reference caches. It is replaced by the native, filterable The 100 table with real blanks. All other cached sortable values agreed with the primary list."),
        ("15 / CHRONOLOGY REPAIR", "The old Timeline contained 100 titles across two side-by-side blocks—not 50 total. It had 61 stale ranks among matching titles, 25 missing years and 24 missing tier cells. It retained Crysis and Metal Gear Solid V: Ground Zeroes, omitted Resident Evil 4 and Resident Evil Requiem, and is now rebuilt exclusively from the 100 primary records."),
        ("16 / METHODOLOGY REPAIR", "Six entries in the old Methodology top-ten side table disagreed with the primary ordering. That duplicate ranking has been removed. The primary source order—not the stale side table or a critic-score sort—is used everywhere."),
        ("17 / REPRODUCIBILITY", "collection.json is the static canonical data for this download. generate_collection.py and requirements.txt reproduce the artifacts from the original workbook; test_collection.py checks them. The manifest and source-audit.json retain exact discrepancies, selected image hashes and extraction evidence. Editing a downloaded workbook does not automatically update the separate JSON."),
        ("18 / SOURCE IDENTITY", SOURCE_NAME + "\nSHA-256: " + audit["sourceSha256"] + "\nThe original attachment is never modified. Generated document timestamps are omitted: no document or critic verification date is invented."),
    ]
    last_methodology_row = 6 + len(notes) * 5
    fill_canvas(methodology, last_methodology_row, 14)
    banner(methodology, "METHODOLOGY & PROVENANCE", "An honest guide to the collection, its numbers, its artwork, and its source.", 14)
    methodology.write_url(3, 4, "internal:'The 100'!A1", nav, "The 100 →")
    label_format = fmt(font_size=10, bold=True, font_color=COLORS["accentInk"], text_wrap=True, valign="top")
    note_format = fmt(font_size=11, font_color=COLORS["secondary"], bg_color=COLORS["panel"], text_wrap=True, valign="top")
    for index, (label, text) in enumerate(notes):
        row = 6 + index * 5
        methodology.merge_range(row, 0, row + 3, 2, label, label_format)
        methodology.merge_range(row, 3, row + 3, 14, text, note_format)
    methodology.freeze_panes(6, 0)
    methodology.print_area(0, 0, last_methodology_row, 14)
    methodology.set_paper(9)
    methodology.fit_to_pages(1, 0)
    methodology.repeat_rows(0, 3)
    methodology.set_h_pagebreaks([36, 66])
    attribution(start, 56, 15)
    attribution(main, 109, 15)
    attribution(chronology, 109, 9)
    attribution(gallery, 309, 18)
    attribution(methodology, last_methodology_row + 2, 14)
    workbook.close()
    omit_generated_document_timestamps(path)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def validate_outputs(collection: dict, audit: dict, selected: list[dict], output: Path) -> dict:
    games = collection["games"]
    require(collection["schemaVersion"] == 1, "Wrong schemaVersion")
    require([game["rank"] for game in games] == list(range(1, 101)), "Ranks are not exactly 1..100")
    for key in ("title", "slug"):
        require(len({game[key] for game in games}) == 100, f"Nonunique {key}")
    originals = {row["rank"]: row for row in audit["primaryRecords"]}
    for game in games:
        require(re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", game["slug"]) is not None, "Nonportable slug")
        require(game["tier"] == ("core" if game["rank"] <= 50 else "essential"), "Incorrect tier")
        require(game["genreTags"] == [item.strip() for item in game["genre"].split("/") if item.strip()], "Invented genre tag")
        require(math.isclose(game["rankIndex"], rank_index(game["rank"]), abs_tol=1e-12), "Incorrect rank index")
        require(game["criticAverage"] == critic_average(game["critics"]), "Incorrect critic average")
        original = originals[game["rank"]]
        row = original["sourceRow"]
        for key, column in (("title", "B"), ("year", "C"), ("studio", "D"), ("genre", "E"), ("rationale", "M")):
            require(game[key] == original["sourceCells"][f"{column}{row}"]["cached"], f"Changed source field {key}")
        for key, column in zip(CRITIC_KEYS, "FGHIJ"):
            require(game["critics"][key] == original["sourceCells"][f"{column}{row}"]["cached"], "Changed source critic")
        note_source = original["sourceCells"][f"L{row}"]
        require(game["sourceNote"] == extract_source_note(note_source["entered"], note_source["cached"]), "Changed source annotation")
        require(game["authorRating"] == original_rating(
            note_source["cached"], note_source["cachedRaw"], f"L{row}", note_source["numberFormat"],
        ), "Changed original author rating or cached display")
    require(len(selected) == sum(game["artwork"] is not None for game in games), "Artwork count mismatch")
    for image in selected:
        game = games[image["rank"] - 1]
        require(image["title"] == game["title"], "Wrong artwork title")
        relative = Path(image["file"])
        require(not relative.is_absolute() and ".." not in relative.parts, "Nonportable artwork path")
        data = (output / relative).read_bytes()
        require(digest(data) == image["sha256"], "Altered extracted image bytes")
        with Image.open(io.BytesIO(data)) as decoded:
            decoded.verify()
        require(game["artwork"]["file"] == image["file"], "JSON artwork path mismatch")
        require(game["artwork"]["source"] == "User-provided workbook", "Wrong artwork attribution")
    workbook_path = output / WORKBOOK_NAME
    with ZipFile(workbook_path) as archive:
        require(archive.testzip() is None, "Workbook ZIP CRC failed")
        xml_count = 0
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(name))
                xml_count += 1
        require(not any("externalLinks/" in name or "vbaProject" in name for name in archive.namelist()), "Unsafe workbook linkage")
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        calc = workbook_xml.find("m:calcPr", NS)
        require(calc is not None and calc.get("calcMode", "auto") == "auto", "Calculation is not automatic")
    formulas = openpyxl.load_workbook(workbook_path, data_only=False)
    cached = openpyxl.load_workbook(workbook_path, data_only=True)
    require(formulas.sheetnames == ["Start Here", "The 100", "Chronology", "Cover Gallery", "Methodology"], "Unexpected sheet list")
    require(formulas["The 100"].freeze_panes == "C7", "Rank/Game/header panes are not frozen")
    require(formulas["The 100"].tables["Play100"].ref == "A6:P106", "Incorrect 100-row native table")
    require(formulas["Chronology"].tables["Play100Chronology"].ref == "A6:J106", "Incomplete chronology table")
    require(tuple(formulas["The 100"].cell(6, index + 1).value for index in range(len(MAIN_HEADERS))) == MAIN_HEADERS, "Wrong headers")
    require(len(formulas["The 100"].conditional_formatting) >= 3, "Missing useful conditional formats")
    require(len(formulas["The 100"].data_validations.dataValidation) == 3, "Missing native-scale validation")
    require(len(formulas["Cover Gallery"]._images) == len(selected), "Incomplete workbook cover gallery")
    score_contrasts = []
    for row in formulas["The 100"].iter_rows(min_row=7, max_row=106, min_col=7, max_col=13):
        for cell in row:
            foreground = cell.font.color.rgb[-6:]
            background = cell.fill.fgColor.rgb[-6:]
            score_contrasts.append(contrast_ratio(foreground, background))
            if cell.column == 12:
                score_contrasts.append(contrast_ratio(foreground, COLORS["accent"]))
    require(min(score_contrasts) >= 7, "Score text must remain high contrast on neutral cells and lime data bars")
    for game in games:
        row = game["rank"] + 6
        sheet = cached["The 100"]
        expected = (
            game["rank"], game["title"], game["tier"].title(), game["year"],
            game["studio"], game["genre"], *game["critics"].values(),
            game["criticAverage"], game["authorRating"]["value"], game["sourceNote"], game["rationale"],
        )
        for column, value in enumerate(expected, start=1):
            actual = sheet.cell(row, column).value
            if is_number(value):
                require(is_number(actual) and math.isclose(actual, value, abs_tol=1e-12), f"Workbook cache mismatch at {row},{column}")
            else:
                require(actual == value, f"Workbook content mismatch at {row},{column}")
        require(formulas["The 100"].cell(row, 12).value == average_formula(row), "Wrong average formula")
        require(formulas["The 100"].cell(row, 13).data_type == "n", "Author rating must be preserved, not recomputed by a formula")
        require(formulas["The 100"].cell(row, 13).number_format == game["authorRating"]["numberFormat"], "Changed source rating number format")
        require(game["authorRating"]["rawValue"] in formulas["The 100"].cell(row, 13).comment.text, "Missing exact original cached rating evidence")
    chronological_games = sorted(games, key=lambda game: (game["year"], game["rank"]))
    for row, game in enumerate(chronological_games, start=7):
        sheet = cached["Chronology"]
        require([sheet.cell(row, column).value for column in (1, 2, 3)] == [game["year"], game["rank"], game["title"]], "Chronology is not complete/source-ordered")
        for column, value, source_column in ((7, game["criticAverage"], "L"), (8, game["authorRating"]["value"], "M")):
            require(math.isclose(sheet.cell(row, column).value, value, abs_tol=1e-12), "Chronology cached metric mismatch")
            require(formulas["Chronology"].cell(row, column).value == main_lookup_formula(source_column, f"B{row}"), "Unstable chronology formula")
    author = collection["collection"]["author"]
    for sheet in formulas:
        links = {cell.hyperlink.target for row in sheet for cell in row if cell.hyperlink and cell.hyperlink.target}
        require({author["githubUrl"], author["linkedinUrl"], author["telegramUrl"]}.issubset(links), f"Missing author footer hyperlinks on {sheet.title}")
        require(author["fullName"] in (sheet.oddFooter.left.text or ""), f"Missing print-footer author on {sheet.title}")
        require(any(author["fullName"] in str(cell.value or "") for row in sheet for cell in row), f"Missing visible author attribution on {sheet.title}")
    error_cells = [
        f"{sheet.title}!{cell.coordinate}"
        for sheet in cached for row in sheet for cell in row if cell.data_type == "e"
    ]
    require(not error_cells, f"Workbook contains cached formula errors: {error_cells}")
    formula_count = sum(cell.data_type == "f" for sheet in formulas for row in sheet for cell in row)
    formulas.close()
    cached.close()
    return {
        "status": "passed",
        "recordCount": 100,
        "uniqueRanksTitlesAndSlugs": True,
        "sourceFieldsAndCriticEntriesPreserved": True,
        "missingCriticValuesAreNullAndBlank": True,
        "criticScoresCompared": 500,
        "missingCriticValueCount": sum(audit["missingCriticCounts"].values()),
        "mainAndChronologyNumericFormulaCachesCompared": 300,
        "originalAuthorRatingsCompared": 100,
        "originalAuthorRatingRawCachesAndFormatsPreserved": True,
        "authorFooterHyperlinksVerifiedOnAllSheets": True,
        "formulaCount": formula_count,
        "formulaExpressionsChecked": True,
        "workbookZipCrcAndXmlValid": True,
        "xmlPartsParsed": xml_count,
        "readableByOpenpyxl": True,
        "nativeTablesContain100RowsEach": True,
        "galleryImagesVerified": len(selected),
        "minimumScoreTextContrastRatio": round(min(score_contrasts), 2),
        "scoreTextContrastAtLeast7To1": True,
        "originalImageBytesPreserved": True,
        "macrosAndExternalWorkbookLinks": False,
        "validationScope": "Library/ZIP/XML checks and direct arithmetic. No Excel/LibreOffice recalculation engine was invoked.",
    }


def make_manifest(collection: dict, audit: dict, selected: list[dict], validation: dict, output: Path) -> dict:
    records = collection["games"]
    differences = {
        "authoritativeOrder": "Preserved all 100 ranks and source titles, studios, years, genres, rationales, and native critic values from AAA Top 50.",
        "primaryAverageHeader": {"sourceCell": "K4", "sourceHeader": "43.2", "replacementHeader": "Critic average /100"},
        "authorRatingPreservation": {
            "sourceHeader": "my rating(based on rank)",
            "replacementHeader": "Leul's original rating /10",
            "meaning": "Leul's original cached ratings, based on his curated rank. Preserve original rounded/text results, values, display, number formats and notes; never substitute the reconstructed curve.",
            "originalFormulaCaveats": audit["rankIndexSourceAnomalies"],
            "visitorRatings": "Separate device-only data. Never seeded from these public author ratings.",
        },
        "sourceCriticAverageDifferences": audit["sourceAverageCacheDifferences"],
        "explicitSourceNotesPreserved": audit["sourceNotes"],
        "blankPrimaryScoresDisplayedAsZeroInSortableRatings": {
            "count": len(audit["sortableRatings"]["blankPrimaryScoresDisplayedAsZero"]),
            "decision": "Use authoritative primary blanks as JSON nulls and genuinely blank Excel critic cells; do not import derivative zeros.",
            "evidence": audit["sortableRatings"]["blankPrimaryScoresDisplayedAsZero"],
        },
        "otherSortableDifferences": audit["sortableRatings"]["otherDifferences"],
        "methodologyAlternateTopTen": {
            "mismatchCount": sum(not item["matches"] for item in audit["methodologyAlternateTopTen"]),
            "evidence": audit["methodologyAlternateTopTen"],
            "decision": "Remove stale duplicate top-ten side table; the primary curated order is the sole order.",
        },
        "timeline": {
            key: value for key, value in audit["timeline"].items() if key != "entries"
        },
        "timelineDecision": "The source actually holds 100 titles in two 50-row blocks. Rebuild one complete 100-record chronology from primary years and ranks; do not reinstate obsolete derivative-only titles.",
        "extraColumnsNO": audit["extraColumnsNO"],
        "separatorRow": {"sourceRow": 55, "decision": "Preserve tier meaning but omit the decorative non-record divider from the native table."},
        "genreTags": "Split only the existing slash-separated source genre; trim separator-adjacent whitespace without adding classifications.",
        "slugs": "NFKD ASCII transliteration, remove apostrophes, replace other punctuation/whitespace with single hyphens; assert uniqueness.",
    }
    caveats = [
        "Scores and source years were not independently reverified. The source provides no reliable critic snapshot date.",
        "Metacritic and Metacritic PC are both included when entered. The normalized average covers entered columns, not independent publications.",
        "Leul's rating is copied from the original cached column L; rank-based provenance is retained. The legacy rankIndex JSON field is a separate internal compatibility value, not a replacement author rating.",
        "An unmarked title has unknown play status. Only 12 explicit (AI – not played) notes are preserved.",
        "91 selected images are low-resolution 120-pixel-high thumbnails; the nine larger poster originals are also modest resolution.",
        "Rank 73 retains source title Hitman: World of Assassination, source year 2016, and its anchored HITMAN III-branded image. This source edition/year naming caveat is disclosed, not silently normalized.",
        "Artwork and visible packaging do not establish platform availability, an exact edition, release date, or redistribution licence.",
        "The source has no document creation timestamp. Automatically generated document created/modified tags are omitted, keeping builds reproducible without inventing a document or critic verification date.",
    ]
    file_paths = [
        output / "collection.json", output / WORKBOOK_NAME, output / "source-audit.json",
        output / "generate_collection.py", output / "test_collection.py", output / "requirements.txt",
        output / "author.json",
        *[output / Path(item["file"]) for item in selected],
    ]
    files = [{
        "path": path.relative_to(output).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": digest(path.read_bytes()),
    } for path in file_paths if path.is_file()]
    return {
        "schemaVersion": 1,
        "title": "Play 100 artifact manifest",
        "workbookPresentation": {
            "identity": "Ink / cream / restrained lime",
            "palette": COLORS,
            "scoreText": "Dark ink on neutral cells and light lime data bars; no lime text on cream.",
            "minimumScoreTextContrastRatio": validation["minimumScoreTextContrastRatio"],
            "headers": "Ink backgrounds with light cream text and restrained lime rules.",
        },
        "source": {
            "displayFilename": SOURCE_NAME, "sha256": audit["sourceSha256"],
            "bytes": audit["sourceSizeBytes"], "authoritativeSheet": PRIMARY_SHEET,
            "originalModified": False, "criticSnapshotDate": None,
            "scope": collection["collection"]["scope"], "rankingBasis": collection["collection"]["rankingBasis"],
            "sheets": audit["sheets"],
        },
        "canonical": {
            "file": "collection.json", "schemaVersion": 1, "gameCount": len(records),
            "coreCount": 50, "essentialCount": 50,
            "sourceYearRange": [min(game["year"] for game in records), max(game["year"] for game in records)],
            "criticColumns": CRITIC_MAX, "missingCounts": audit["missingCriticCounts"],
            "availableCriticEntryCount": 500 - sum(audit["missingCriticCounts"].values()),
            "explicitSourceNoteCount": len(audit["sourceNotes"]),
            "criticAverage": "Arithmetic mean of the available columns after IGN and GameSpot × 10; missing => null; all missing => null.",
            "rankIndex": "10 - (rank - 1) * 3 / 99; a derived index, not an independent review.",
            "authorRating": "Original numeric value plus exact raw cached text, displayed source value, source cell and number format; all100 are preserved.",
            "author": collection["collection"]["author"],
            "precision": "JSON retains full finite Python floating-point precision; workbook displays scores to one decimal without rounding stored values.",
            "artworkCount": len(selected),
        },
        "differencesAndCleaningDecisions": differences,
        "artwork": {
            "source": "User-provided workbook",
            "sourceEmbeddedImageCount": len(audit["embeddedImages"]),
            "rowThumbnails": 100, "largerPosters": 9,
            "selectedCount": len(selected),
            "selectedFrom": dict(Counter(item["sourceSheet"] for item in selected)),
            "mapping": "Resolve drawing relationships and exact anchor-cell titles; do not assume drawing order equals curated order. Review contact sheets against those labels.",
            "allPosterTitlesMatchPrimary": True,
            "externalArtworkRequests": 0,
            "imageBytesModified": False,
            "selected": selected,
            "allSourceImagesEvidence": "source-audit.json#embeddedImages",
            "visualReviewEvidence": audit["contactSheets"],
        },
        "uncertaintiesAndLimitations": caveats,
        "libraryReadWarnings": audit["libraryReadWarnings"],
        "libraryReadWarningHandling": audit["libraryReadWarningHandling"],
        "reproduce": {
            "generator": "generate_collection.py", "tests": "test_collection.py", "requirements": "requirements.txt",
            "pythonVersion": sys.version.split()[0],
            "libraryVersions": {"openpyxl": openpyxl.__version__, "XlsxWriter": xlsxwriter.__version__, "Pillow": PIL.__version__},
            "windowsCommands": [
                "python -m venv .venv",
                r".\.venv\Scripts\python.exe -m pip install -r requirements.txt",
                r'.\.venv\Scripts\python.exe -B .\generate_collection.py --source "PATH_TO_ORIGINAL.xlsx" --output .',
                r".\.venv\Scripts\python.exe -B -m unittest -q test_collection",
            ],
            "sourcePinnedArtworkReview": REVIEWED_SOURCE_SHA256,
            "differentSourceHandling": "Use --inspect-only to inspect a different workbook. Generation refuses unreviewed source bytes rather than silently assigning unverified artwork.",
            "networkRequiredDuringGeneration": False,
            "outputPathDefaultsToScriptDirectory": True,
            "copyForWebsite": ["collection.json", "assets", WORKBOOK_NAME],
            "shareableWorkbook": "The .xlsx embeds its own images, tables, cached values and methodology. No sibling files are required to open it.",
            "excludeFromDistribution": [".venv", ".package-work", "__pycache__", "previews"],
        },
        "validation": validation,
        "files": files,
        "manifestSelfHash": "Not included to avoid a recursive self-hash.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--inspect-only", action="store_true")
    parser.add_argument("--author-config", type=Path)
    args = parser.parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    require(source != (output / WORKBOOK_NAME).resolve(), "The output workbook would overwrite the source.")
    output.mkdir(parents=True, exist_ok=True)
    collection, audit, candidates = read_source(source)
    collection["collection"]["author"] = read_author_config(args.author_config)
    collection["collection"]["authorRatingsAreOriginal"] = True
    if not args.inspect_only:
        require(
            audit["sourceSha256"] == REVIEWED_SOURCE_SHA256,
            "The source bytes differ from the visually reviewed workbook; inspect with --inspect-only before updating the review pin.",
        )
    audit["contactSheets"] = make_contact_sheets(candidates, output)
    write_json(output / "author.json", collection["collection"]["author"])
    write_json(output / "source-audit.json", audit)
    if not args.inspect_only:
        selected = select_artwork(collection, candidates, output)
        write_json(output / "collection.json", collection)
        build_workbook(collection, audit, selected, output)
        validation = validate_outputs(collection, audit, selected, output)
        require(digest(source.read_bytes()) == audit["sourceSha256"], "Original source changed during generation")
        validation["sourceSha256UnchangedAfterGeneration"] = True
        manifest = make_manifest(collection, audit, selected, validation, output)
        write_json(output / "artifact-manifest.json", manifest)
    print(json.dumps({
        "records": len(collection["games"]), "embeddedImages": len(candidates),
        "notes": len(audit["sourceNotes"]),
        "timelineRecords": audit["timeline"]["recordCount"],
        "timelineStaleRanks": audit["timeline"]["rankDisagreementCount"],
        "timelineMissingYears": audit["timeline"]["missingYearCount"],
        "timelineMissingTiers": audit["timeline"]["missingTierCount"],
        "sourceSha256": audit["sourceSha256"],
        "output": str(output),
        "mode": "inspect-only" if args.inspect_only else "generated-and-validated",
    }, indent=2))


if __name__ == "__main__":
    main()
