"""Focused checks for Play 100. Run: python -B -m unittest -q test_collection."""

import io
import json
import math
import unittest
from fractions import Fraction
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import xlsxwriter

from generate_collection import (
    COLORS,
    CRITIC_KEYS,
    REVIEWED_SOURCE_SHA256,
    SOURCE_NAME,
    WORKBOOK_NAME,
    average_formula,
    critic_average,
    contrast_ratio,
    digest,
    extract_source_note,
    main_link_formula,
    original_rating,
    rank_index,
    slugify,
    validate_outputs,
)


ROOT = Path(__file__).resolve().parent


class ArithmeticTests(unittest.TestCase):
    def test_no_scores_is_none_not_zero(self):
        self.assertIsNone(critic_average(dict.fromkeys(CRITIC_KEYS)))

    def test_each_native_scale_independently(self):
        for key, entered, normalized in (
            ("metacritic", 91, 91), ("metacriticPc", 89, 89),
            ("ign", 9.3, 93), ("gamespot", 8.5, 85), ("pcGamer", 84, 84),
        ):
            with self.subTest(critic=key):
                values = dict.fromkeys(CRITIC_KEYS)
                values[key] = entered
                self.assertAlmostEqual(critic_average(values), normalized)

    def test_both_metacritic_columns_count(self):
        scores = dict(zip(CRITIC_KEYS, (97, 93, 10, 9, None)))
        self.assertEqual(critic_average(scores), 95)

    def test_missing_values_are_not_counted(self):
        scores = dict(zip(CRITIC_KEYS, (95, None, 10, 8, None)))
        self.assertAlmostEqual(critic_average(scores), (95 + 100 + 80) / 3)

    def test_real_zero_remains_a_value(self):
        scores = dict.fromkeys(CRITIC_KEYS)
        scores["ign"] = 0
        self.assertEqual(critic_average(scores), 0)
        scores["metacritic"] = 100
        self.assertEqual(critic_average(scores), 50)

    def test_rank_index_exact_endpoints_and_uniform_steps(self):
        self.assertEqual(rank_index(1), 10)
        self.assertEqual(rank_index(100), 7)
        for rank in range(1, 101):
            expected = float(Fraction(331 - rank, 33))
            self.assertAlmostEqual(rank_index(rank), expected, places=12)
        self.assertTrue(all(rank_index(rank) > rank_index(rank + 1) for rank in range(1, 100)))

    def test_notes_do_not_become_reviews_or_play_status(self):
        formula = '=TEXT(10-((A7-1)*(3/99)),"0.0")&" (AI – not played)"'
        self.assertEqual(extract_source_note(formula, "9.9 (AI – not played)"), "(AI – not played)")
        self.assertEqual(extract_source_note(formula, None), "(AI – not played)")
        self.assertIsNone(extract_source_note('=TEXT(10-(5*(3/99)),"0.0")', "9.8"))
        self.assertIsNone(extract_source_note("=10-((A5-1)*(3/99))", 10))
        self.assertEqual(extract_source_note("(AI – not played)", "(AI – not played)"), "(AI – not played)")

    def test_slug_transliteration_and_punctuation(self):
        self.assertEqual(slugify("God of War Ragnarök"), "god-of-war-ragnarok")
        self.assertEqual(slugify("Baldur's Gate 3"), "baldurs-gate-3")
        self.assertEqual(slugify("L.A. Noire"), "l-a-noire")

    def test_blank_and_zero_formula_caches_roundtrip(self):
        buffer = io.BytesIO()
        with xlsxwriter.Workbook(buffer, {"in_memory": True}) as workbook:
            sheet = workbook.add_worksheet()
            sheet.write_formula("L7", average_formula(7), None, "")
            sheet.write_number("I8", 0)
            sheet.write_formula("L8", average_formula(8), None, 0)
        contents = buffer.getvalue()
        cached = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(contents), data_only=False)
        self.assertIn(cached.active["L7"].value, (None, ""))
        self.assertEqual(cached.active["L8"].value, 0)
        self.assertEqual(formulas.active["L7"].value, average_formula(7))
        self.assertIn("COUNT(G7:K7)=0", average_formula(7))
        cached.close()
        formulas.close()

    def test_hyperlinks_lookup_rank_after_sort(self):
        formula = main_link_formula("B12")
        self.assertIn("MATCH(B12,'The 100'!$A$7:$A$106,0)", formula)
        self.assertIn("+6", formula)

    def test_lime_palette_retains_readable_text(self):
        for background in ("background", "panel", "alternate", "accent"):
            with self.subTest(background=background):
                self.assertGreaterEqual(contrast_ratio(COLORS["text"], COLORS[background]), 7)
        self.assertGreaterEqual(contrast_ratio(COLORS["accent"], COLORS["ink"]), 7)
        self.assertGreaterEqual(contrast_ratio(COLORS["accentInk"], COLORS["accentPanel"]), 7)
        for background in ("background", "panel", "alternate"):
            self.assertGreaterEqual(contrast_ratio(COLORS["muted"], COLORS[background]), 4.5)


class ArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.collection = json.loads((ROOT / "collection.json").read_text(encoding="utf-8"))
        cls.audit = json.loads((ROOT / "source-audit.json").read_text(encoding="utf-8"))
        cls.manifest = json.loads((ROOT / "artifact-manifest.json").read_text(encoding="utf-8"))

    def test_full_cross_artifact_validation(self):
        result = validate_outputs(
            self.collection, self.audit, self.manifest["artwork"]["selected"], ROOT,
        )
        self.assertEqual(result["status"], "passed")
        self.assertEqual(result["mainAndChronologyNumericFormulaCachesCompared"], 300)
        self.assertEqual(result["originalAuthorRatingsCompared"], 100)
        self.assertTrue(result["authorFooterHyperlinksVerifiedOnAllSheets"])

    def test_primary_order_and_manual_edits_preserved(self):
        games = self.collection["games"]
        self.assertEqual([game["title"] for game in games[:6]], [
            "Red Dead Redemption 2", "Mass Effect 2", "The Witcher 3: Wild Hunt",
            "God of War", "The Last of Us", "Grand Theft Auto V",
        ])
        self.assertEqual((games[26]["title"], games[26]["year"]), ("Resident Evil 4", 2023))
        self.assertEqual((games[40]["title"], games[40]["year"]), ("Resident Evil Requiem", 2026))
        self.assertEqual(games[-1]["title"], "L.A. Noire")
        self.assertNotIn("Crysis", {game["title"] for game in games})
        self.assertNotIn("Metal Gear Solid V: Ground Zeroes", {game["title"] for game in games})

    def test_schema_contract_and_provenance(self):
        expected_keys = {
            "rank", "slug", "title", "year", "studio", "genre", "genreTags", "tier",
            "critics", "criticAverage", "rankIndex", "authorRating", "rationale", "sourceNote", "artwork",
        }
        self.assertEqual(self.collection["schemaVersion"], 1)
        self.assertEqual(self.collection["collection"]["sourceFile"], SOURCE_NAME)
        self.assertTrue(self.collection["collection"]["criticScoresAreSnapshot"])
        for game in self.collection["games"]:
            self.assertEqual(set(game), expected_keys)
            self.assertEqual(set(game["critics"]), set(CRITIC_KEYS))
            self.assertIsInstance(game["year"], int)
            self.assertNotIn("played", game)
            for score in game["critics"].values():
                self.assertTrue(score is None or math.isfinite(score))

    def test_independent_fraction_arithmetic(self):
        for game in self.collection["games"]:
            inputs = [
                Fraction(str(value)) * (10 if key in ("ign", "gamespot") else 1)
                for key, value in game["critics"].items() if value is not None
            ]
            expected = float(sum(inputs, Fraction(0)) / len(inputs)) if inputs else None
            self.assertAlmostEqual(game["criticAverage"], expected, places=12)
            self.assertAlmostEqual(game["rankIndex"], float(Fraction(331 - game["rank"], 33)), places=12)

    def test_missing_values_and_explicit_notes(self):
        games = self.collection["games"]
        expected_missing = {"metacritic": 0, "metacriticPc": 16, "ign": 16, "gamespot": 0, "pcGamer": 24}
        self.assertEqual(
            {key: sum(game["critics"][key] is None for game in games) for key in CRITIC_KEYS},
            expected_missing,
        )
        self.assertEqual(
            [game["rank"] for game in games if game["sourceNote"]],
            [3, 11, 12, 15, 22, 24, 28, 30, 34, 38, 48, 49],
        )
        self.assertTrue(all(game["sourceNote"] in (None, "(AI – not played)") for game in games))

    def test_actual_derivative_discrepancies_documented(self):
        differences = self.manifest["differencesAndCleaningDecisions"]
        self.assertEqual(differences["methodologyAlternateTopTen"]["mismatchCount"], 6)
        self.assertEqual(differences["blankPrimaryScoresDisplayedAsZeroInSortableRatings"]["count"], 56)
        self.assertEqual(differences["timeline"]["recordCount"], 100)
        self.assertEqual(differences["timeline"]["rankDisagreementCount"], 61)
        self.assertEqual(differences["timeline"]["missingYearCount"], 25)
        self.assertEqual(differences["timeline"]["missingTierCount"], 24)
        self.assertEqual(differences["authorRatingPreservation"]["originalFormulaCaveats"][0]["rank"], 7)
        self.assertEqual(differences["extraColumnsNO"], {})
        self.assertEqual(differences["sourceCriticAverageDifferences"], [])

    def test_source_identity_and_artwork_count(self):
        self.assertEqual(self.audit["sourceSha256"], REVIEWED_SOURCE_SHA256)
        self.assertFalse(self.manifest["source"]["originalModified"])
        artwork = self.manifest["artwork"]
        self.assertEqual(artwork["sourceEmbeddedImageCount"], 109)
        self.assertEqual(artwork["selectedCount"], 100)
        self.assertEqual(artwork["selectedFrom"], {"Posters": 9, "AAA Top 50": 91})
        self.assertEqual(
            sorted(item["rank"] for item in artwork["selected"] if item["sourceSheet"] == "Posters"),
            [1, 2, 3, 4, 5, 6, 9, 13, 15],
        )

    def test_exact_original_author_ratings_are_preserved_not_reconstructed(self):
        originals = {row["rank"]: row for row in self.audit["primaryRecords"]}
        for game in self.collection["games"]:
            source = originals[game["rank"]]
            cell = f'L{source["sourceRow"]}'
            original = source["sourceCells"][cell]
            expected = original_rating(original["cached"], original["cachedRaw"], cell, original["numberFormat"])
            self.assertEqual(game["authorRating"], expected)
        self.assertEqual(self.collection["games"][1]["authorRating"]["rawValue"], "9.9696969696969688")
        self.assertEqual(self.collection["games"][2]["authorRating"]["value"], 9.9)
        self.assertEqual(self.collection["games"][2]["authorRating"]["rawValue"], "9.9 (AI – not played)")
        self.assertEqual(self.collection["games"][6]["authorRating"]["value"], 9.8)
        self.assertEqual(self.collection["games"][6]["authorRating"]["rawValue"], "9.8")
        self.assertNotEqual(self.collection["games"][2]["authorRating"]["value"], self.collection["games"][2]["rankIndex"])
        self.assertNotEqual(self.collection["games"][6]["authorRating"]["value"], self.collection["games"][6]["rankIndex"])

    def test_shared_author_identity_and_hyperlinked_sheet_footers(self):
        author = self.collection["collection"]["author"]
        self.assertEqual(author["fullName"], "Leul Tewodros Agonafer")
        self.assertEqual(author["githubUrl"], "https://github.com/LeulTew/play-100")
        self.assertEqual(author["linkedinUrl"], "https://www.linkedin.com/in/leul-t-agonafer-861bb3336/")
        self.assertEqual(author["telegramUrl"], "https://t.me/fabbin")
        workbook = openpyxl.load_workbook(ROOT / WORKBOOK_NAME, data_only=False)
        self.assertEqual(workbook.properties.creator, author["fullName"])
        self.assertEqual(workbook["The 100"]["M6"].value, "Leul's original rating /10")
        for sheet in workbook:
            links = {cell.hyperlink.target for row in sheet for cell in row if cell.hyperlink and cell.hyperlink.target}
            self.assertTrue({author["githubUrl"], author["linkedinUrl"], author["telegramUrl"]}.issubset(links))
            self.assertIn(author["fullName"], sheet.oddFooter.left.text)
        self.assertEqual(workbook["The 100"].tables["Play100"].ref, "A6:P106")
        self.assertEqual(workbook["Chronology"].tables["Play100Chronology"].ref, "A6:J106")
        workbook.close()

    def test_generated_package_does_not_invent_dates(self):
        self.assertIsNone(self.audit["sourceDocumentCreated"])
        self.assertIsNone(self.manifest["source"]["criticSnapshotDate"])
        with ZipFile(ROOT / WORKBOOK_NAME) as archive:
            core = ET.fromstring(archive.read("docProps/core.xml"))
            self.assertIsNone(core.find("{http://purl.org/dc/terms/}created"))
            self.assertIsNone(core.find("{http://purl.org/dc/terms/}modified"))

    def test_workbook_has_ink_headers_and_neutral_score_cells(self):
        workbook = openpyxl.load_workbook(ROOT / WORKBOOK_NAME, data_only=False)
        header = workbook["The 100"]["G6"]
        score = workbook["The 100"]["G7"]
        self.assertEqual(header.fill.fgColor.rgb[-6:], COLORS["ink"].lstrip("#"))
        self.assertEqual(header.font.color.rgb[-6:], COLORS["onInk"].lstrip("#"))
        self.assertEqual(score.font.color.rgb[-6:], COLORS["text"].lstrip("#"))
        self.assertEqual(score.fill.fgColor.rgb[-6:], COLORS["panel"].lstrip("#"))
        self.assertGreaterEqual(self.manifest["validation"]["minimumScoreTextContrastRatio"], 7)
        self.assertEqual(self.manifest["workbookPresentation"]["identity"], "Ink / cream / restrained lime")
        workbook.close()

    def test_file_checksums_and_no_external_links(self):
        for item in self.manifest["files"]:
            path = ROOT / Path(item["path"])
            self.assertEqual(path.stat().st_size, item["bytes"], str(path))
            self.assertEqual(digest(path.read_bytes()), item["sha256"], str(path))
        with ZipFile(ROOT / WORKBOOK_NAME) as archive:
            self.assertIsNone(archive.testzip())
            self.assertFalse(any("externalLinks/" in name or "vbaProject" in name for name in archive.namelist()))


if __name__ == "__main__":
    unittest.main()
